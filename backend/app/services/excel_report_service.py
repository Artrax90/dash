import io
import time
from datetime import datetime, timezone, timedelta
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference, Series
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

def _clean_val(val: Any) -> Any:
    """Sanitize strings to eliminate illegal characters prohibited in XML / OpenPyXL."""
    if val is None:
        return ''
    if isinstance(val, (int, float, bool)):
        return val
    s = str(val)
    # Remove control characters that cause openpyxl.utils.exceptions.IllegalCharacterError
    return ILLEGAL_CHARACTERS_RE.sub('', s).strip()

def _set_cell(ws, row: int, column: int, value: Any, alignment: Optional[Alignment] = None):
    """Safely set a cell value with illegal characters filtered out and optional alignment."""
    cell = ws.cell(row=row, column=column, value=_clean_val(value))
    if alignment is not None:
        cell.alignment = alignment
    return cell

def _format_dt_local(val: Any, fmt: str = '%d.%m.%Y %H:%M:%S') -> str:
    """Format any date/time or timestamp value into local system timezone without UTC confusion."""
    if val is None or val == '':
        return '—'
    try:
        if isinstance(val, (int, float)):
            return datetime.fromtimestamp(val).astimezone().strftime(fmt)
        if isinstance(val, datetime):
            if val.tzinfo is None:
                val = val.replace(tzinfo=timezone.utc)
            return val.astimezone().strftime(fmt)
        s = str(val).strip()
        if not s:
            return '—'
        # Try ISO format
        clean_s = s.replace('Z', '+00:00')
        try:
            dt = datetime.fromisoformat(clean_s)
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            return dt.astimezone().strftime(fmt)
        except Exception:
            pass
        # Try float timestamp string
        try:
            f = float(s)
            return datetime.fromtimestamp(f).astimezone().strftime(fmt)
        except Exception:
            pass
        return s
    except Exception:
        return str(val)

def _pt_time(p: Dict[str, Any]) -> float:
    t = p.get('time')
    if isinstance(t, (int, float)) and t > 0:
        return float(t)
    iso = p.get('timestamp')
    if iso:
        try:
            return datetime.fromisoformat(str(iso).replace('Z', '+00:00')).timestamp()
        except Exception:
            pass
    return 0.0

def _apply_header_style(cell, text: str, bg_color: str = '0F172A', font_color: str = 'FFFFFF'):
    cell.value = _clean_val(text)
    cell.font = Font(name='Segoe UI', size=11, bold=True, color=font_color)
    cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def _apply_kpi_card(ws, start_row: int, start_col: int, title: str, value: str, subtext: str, color_hex: str):
    fill_bg = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    c_title = ws.cell(row=start_row, column=start_col)
    c_title.value = _clean_val(title.upper())
    c_title.font = Font(name='Segoe UI', size=9, bold=True, color='64748B')
    c_title.alignment = Alignment(horizontal='left', vertical='center')
    
    c_val = ws.cell(row=start_row + 1, column=start_col)
    c_val.value = _clean_val(value)
    c_val.font = Font(name='Segoe UI', size=18, bold=True, color=color_hex)
    c_val.alignment = Alignment(horizontal='left', vertical='center')
    
    c_sub = ws.cell(row=start_row + 2, column=start_col)
    c_sub.value = _clean_val(subtext)
    c_sub.font = Font(name='Segoe UI', size=9, italic=True, color='94A3B8')
    c_sub.alignment = Alignment(horizontal='left', vertical='center')

    for r in range(start_row, start_row + 3):
        for c in range(start_col, start_col + 2):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill_bg
            cell.border = thin_border
    
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col + 1)
    ws.merge_cells(start_row=start_row + 1, start_column=start_col, end_row=start_row + 1, end_column=start_col + 1)
    ws.merge_cells(start_row=start_row + 2, start_column=start_col, end_row=start_row + 2, end_column=start_col + 1)

def generate_monitoring_excel_report(
    period_type: str,
    scope_title: str,
    devices: List[Dict[str, Any]],
    telemetry_points: List[Dict[str, Any]],
    power_events: List[Dict[str, Any]],
    alerts: List[Dict[str, Any]],
    start_ts: Optional[float] = None,
    end_ts: Optional[float] = None
) -> bytes:
    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = 'Сводка и Графики'
    ws_summary.views.sheetView[0].showGridLines = True

    # Build device lookup map for resolving PC names across all sheets
    dev_map: Dict[str, str] = {}
    for d in devices:
        d_id = str(d.get('id', '')).strip()
        name = str(d.get('name', '')).strip() or d_id
        if d_id:
            dev_map[d_id] = name
            dev_map[d_id.upper()] = name
            dev_map[d_id.lower()] = name
        h_name = str(d.get('hostname', '')).strip()
        if h_name:
            dev_map[h_name] = name
            dev_map[h_name.upper()] = name
            dev_map[h_name.lower()] = name

    def _get_dev_name(dev_id: Any, explicit_name: Optional[str] = None) -> str:
        if explicit_name and str(explicit_name).strip():
            return str(explicit_name).strip()
        s_id = str(dev_id or '').strip()
        if not s_id:
            return '—'
        return dev_map.get(s_id, dev_map.get(s_id.upper(), dev_map.get(s_id.lower(), s_id)))

    # 1. Header Banner
    ws_summary.row_dimensions[1].height = 32
    ws_summary.row_dimensions[2].height = 24
    ws_summary.merge_cells('A1:K1')
    c1 = ws_summary['A1']
    c1.value = _clean_val('WORKSTATION MANAGER · ОТЧЕТ МОНИТОРИНГА И ТЕЛЕМЕТРИИ')
    c1.font = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
    c1.fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
    c1.alignment = Alignment(horizontal='center', vertical='center')
    
    period_names = {
        'hourly_24h': 'По часам (последние 24 часа)',
        'daily_7d': 'По дням (последние 7 дней)',
        'daily_30d': 'По дням (последние 30 дней)',
        'weekly_12w': 'По неделям (последние 12 недель)',
        'custom': 'Пользовательский период'
    }
    period_display = period_names.get(period_type, period_type)
    now_local_str = datetime.now().astimezone().strftime('%d.%m.%Y %H:%M')
    
    ws_summary.merge_cells('A2:K2')
    c2 = ws_summary['A2']
    c2.value = _clean_val(f'Зона охвата: {scope_title}   |   Период: {period_display}   |   Сформирован: {now_local_str}')
    c2.font = Font(name='Segoe UI', size=10, italic=False, color='334155')
    c2.fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    c2.alignment = Alignment(horizontal='center', vertical='center')

    total_devs = len(devices)
    online_devs = len([d for d in devices if str(d.get('power_status', d.get('powerStatus', ''))).lower() in ['on', 'booting']])
    offline_devs = total_devs - online_devs

    online_pts = [p for p in telemetry_points if p.get('isOnline', True)]
    avg_cpu = round(sum(p.get('cpu', 0) for p in online_pts) / len(online_pts)) if online_pts else 0
    max_cpu = max([p.get('cpu', 0) for p in online_pts]) if online_pts else 0
    avg_ram = round(sum(p.get('ram', 0) for p in online_pts) / len(online_pts)) if online_pts else 0
    max_ram = max([p.get('ram', 0) for p in online_pts]) if online_pts else 0
    avg_disk = round(sum(p.get('disk', 0) for p in online_pts) / len(online_pts)) if online_pts else 0

    _apply_kpi_card(ws_summary, start_row=4, start_col=1, title='Станции онлайн', value=f'{online_devs} / {total_devs}', subtext=f'{offline_devs} офлайн', color_hex='059669')
    _apply_kpi_card(ws_summary, start_row=4, start_col=3, title='Средняя утилизация ЦП', value=f'{avg_cpu}%', subtext=f'Пик: {max_cpu}%', color_hex='2563EB')
    _apply_kpi_card(ws_summary, start_row=4, start_col=5, title='Использование памяти', value=f'{avg_ram}%', subtext=f'Пик: {max_ram}%', color_hex='7C3AED')
    _apply_kpi_card(ws_summary, start_row=4, start_col=7, title='Средняя занятость дисков', value=f'{avg_disk}%', subtext='По парку станций', color_hex='0891B2')
    _apply_kpi_card(ws_summary, start_row=4, start_col=9, title='События и инциденты', value=f'{len(power_events)} соб. / {len(alerts)} ал.', subtext='За выбранный период', color_hex='D97706')

    # 2. Time Buckets Progression for Charts & Summary Table
    now_ts = time.time()
    period_span_map = {
        '1h': 3600,
        '6h': 21600,
        '24h': 86400,
        'hourly_24h': 86400,
        '7d': 7 * 86400,
        'daily_7d': 7 * 86400,
        '30d': 30 * 86400,
        'daily_30d': 30 * 86400,
        '12w': 84 * 86400,
        'weekly_12w': 84 * 86400,
    }

    if end_ts is None:
        end_ts = now_ts
    if start_ts is None:
        fallback_span = period_span_map.get(period_type, 86400)
        start_ts = end_ts - fallback_span

    if end_ts <= start_ts or (end_ts - start_ts) < 60:
        fallback_span = period_span_map.get(period_type, 86400)
        start_ts = end_ts - fallback_span

    bucket_count = 12
    span = end_ts - start_ts
    step = span / bucket_count

    chart_table_start_row = 9
    ws_summary.merge_cells('A8:E8')
    t_title = ws_summary['A8']
    t_title.value = _clean_val('Сводная динамика нагрузки для диаграмм')
    t_title.font = Font(name='Segoe UI', size=11, bold=True, color='1E293B')
    t_title.alignment = Alignment(horizontal='left', vertical='center')

    chart_headers = ['Интервал времени', 'Загрузка ЦП (%)', 'Память ОЗУ (%)', 'Диск (%)', 'ПК в сети (шт)']
    for col_idx, h_name in enumerate(chart_headers, start=1):
        _apply_header_style(ws_summary.cell(row=chart_table_start_row, column=col_idx), h_name, bg_color='1E293B')

    current_row = chart_table_start_row + 1
    for i in range(bucket_count):
        b_start = start_ts + (i * step)
        b_end = b_start + step
        b_pts = [p for p in telemetry_points if b_start <= _pt_time(p) <= b_end]
        b_online = [p for p in b_pts if p.get('isOnline', True)]
        
        c_val = round(sum(p.get('cpu', 0) for p in b_online) / len(b_online)) if b_online else 0
        r_val = round(sum(p.get('ram', 0) for p in b_online) / len(b_online)) if b_online else 0
        d_val = round(sum(p.get('disk', 0) for p in b_online) / len(b_online)) if b_online else 0
        cnt_val = len(set(p.get('deviceId', 'PC') for p in b_online)) if b_online else (online_devs if i == bucket_count - 1 else 0)

        dt_label = datetime.fromtimestamp(b_start).astimezone().strftime('%d.%m %H:%M')
        _set_cell(ws_summary, row=current_row, column=1, value=dt_label, alignment=Alignment(horizontal='center'))
        _set_cell(ws_summary, row=current_row, column=2, value=c_val, alignment=Alignment(horizontal='center'))
        _set_cell(ws_summary, row=current_row, column=3, value=r_val, alignment=Alignment(horizontal='center'))
        _set_cell(ws_summary, row=current_row, column=4, value=d_val, alignment=Alignment(horizontal='center'))
        _set_cell(ws_summary, row=current_row, column=5, value=cnt_val, alignment=Alignment(horizontal='center'))
        current_row += 1

    chart_table_end_row = current_row - 1

    chart = LineChart()
    chart.title = 'Динамика загрузки ЦП, Памяти и Дисков (%)'
    chart.style = 13
    chart.y_axis.title = 'Процент утилизации (%)'
    chart.x_axis.title = 'Временная шкала'
    chart.width = 22
    chart.height = 12

    data_ref = Reference(ws_summary, min_col=2, min_row=chart_table_start_row, max_col=4, max_row=chart_table_end_row)
    cats_ref = Reference(ws_summary, min_col=1, min_row=chart_table_start_row + 1, max_row=chart_table_end_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)

    colors = ['2563EB', '7C3AED', '0891B2']
    for idx, s in enumerate(chart.series):
        s.graphicalProperties.line.width = 25000
        if idx < len(colors):
            s.graphicalProperties.line.solidFill = colors[idx]

    ws_summary.add_chart(chart, 'G8')

    chart_active = LineChart()
    chart_active.title = 'Количество активных станций онлайн (шт)'
    chart_active.style = 10
    chart_active.y_axis.title = 'ПК онлайн'
    chart_active.width = 22
    chart_active.height = 8

    data_active_ref = Reference(ws_summary, min_col=5, min_row=chart_table_start_row, max_col=5, max_row=chart_table_end_row)
    chart_active.add_data(data_active_ref, titles_from_data=True)
    chart_active.set_categories(cats_ref)
    if chart_active.series:
        chart_active.series[0].graphicalProperties.line.solidFill = '059669'
        chart_active.series[0].graphicalProperties.line.width = 25000

    ws_summary.add_chart(chart_active, 'G22')

    # Sheet 2: Raw Telemetry
    ws_telem = wb.create_sheet(title='Телеметрия')
    ws_telem.views.sheetView[0].showGridLines = True
    telem_headers = ['Дата и Время', 'ID ПК', 'Имя ПК', 'Загрузка ЦП (%)', 'Память ОЗУ (%)', 'Диск (%)', 'Статус питания']
    for idx, h in enumerate(telem_headers, start=1):
        _apply_header_style(ws_telem.cell(row=1, column=idx), h, bg_color='334155')

    sorted_pts = sorted(telemetry_points, key=lambda p: _pt_time(p))
    for r_idx, pt in enumerate(sorted_pts, start=2):
        dev_id = pt.get('deviceId', '')
        dev_name = _get_dev_name(dev_id, pt.get('deviceName'))
        dt_str = _format_dt_local(pt.get('timestamp') or pt.get('time'))
        status_str = 'В сети' if pt.get('isOnline', True) else 'Выключен'
        
        _set_cell(ws_telem, row=r_idx, column=1, value=dt_str, alignment=Alignment(horizontal='center'))
        _set_cell(ws_telem, row=r_idx, column=2, value=dev_id, alignment=Alignment(horizontal='center'))
        _set_cell(ws_telem, row=r_idx, column=3, value=dev_name, alignment=Alignment(horizontal='left'))
        _set_cell(ws_telem, row=r_idx, column=4, value=pt.get('cpu', 0), alignment=Alignment(horizontal='center'))
        _set_cell(ws_telem, row=r_idx, column=5, value=pt.get('ram', 0), alignment=Alignment(horizontal='center'))
        _set_cell(ws_telem, row=r_idx, column=6, value=pt.get('disk', 0), alignment=Alignment(horizontal='center'))
        _set_cell(ws_telem, row=r_idx, column=7, value=status_str, alignment=Alignment(horizontal='center'))

    # Sheet 3: Power Events
    ws_power = wb.create_sheet(title='События питания')
    ws_power.views.sheetView[0].showGridLines = True
    power_headers = ['ID события', 'Дата и Время', 'Действие', 'Название события', 'ID ПК', 'Имя ПК', 'Инициатор', 'Источник', 'Статус', 'Детали']
    for idx, h in enumerate(power_headers, start=1):
        _apply_header_style(ws_power.cell(row=1, column=idx), h, bg_color='1E3A8A')

    for r_idx, ev in enumerate(power_events, start=2):
        dev_id = ev.get('deviceId', '')
        dev_name = _get_dev_name(dev_id, ev.get('deviceName'))
        dt_str = _format_dt_local(ev.get('timestamp'))

        _set_cell(ws_power, row=r_idx, column=1, value=ev.get('id', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_power, row=r_idx, column=2, value=dt_str, alignment=Alignment(horizontal='center'))
        _set_cell(ws_power, row=r_idx, column=3, value=ev.get('action', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_power, row=r_idx, column=4, value=ev.get('title', ''), alignment=Alignment(horizontal='left'))
        _set_cell(ws_power, row=r_idx, column=5, value=dev_id, alignment=Alignment(horizontal='center'))
        _set_cell(ws_power, row=r_idx, column=6, value=dev_name, alignment=Alignment(horizontal='left'))
        _set_cell(ws_power, row=r_idx, column=7, value=ev.get('initiator', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_power, row=r_idx, column=8, value=ev.get('source', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_power, row=r_idx, column=9, value=ev.get('status', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_power, row=r_idx, column=10, value=ev.get('details', ''), alignment=Alignment(horizontal='left'))

    # Sheet 4: Alerts
    ws_alerts = wb.create_sheet(title='Алерты и Инциденты')
    ws_alerts.views.sheetView[0].showGridLines = True
    alert_headers = ['ID алерта', 'Дата и Время', 'ID ПК', 'Имя ПК', 'Критичность', 'Категория', 'Описание проблемы']
    for idx, h in enumerate(alert_headers, start=1):
        _apply_header_style(ws_alerts.cell(row=1, column=idx), h, bg_color='991B1B')

    for r_idx, alt in enumerate(alerts, start=2):
        dev_id = alt.get('deviceId', alt.get('device_id', ''))
        dev_name = _get_dev_name(dev_id, alt.get('deviceName'))
        dt_str = _format_dt_local(alt.get('timestamp') or alt.get('created_at'))

        _set_cell(ws_alerts, row=r_idx, column=1, value=alt.get('id', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_alerts, row=r_idx, column=2, value=dt_str, alignment=Alignment(horizontal='center'))
        _set_cell(ws_alerts, row=r_idx, column=3, value=dev_id, alignment=Alignment(horizontal='center'))
        _set_cell(ws_alerts, row=r_idx, column=4, value=dev_name, alignment=Alignment(horizontal='left'))
        _set_cell(ws_alerts, row=r_idx, column=5, value=alt.get('severity', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_alerts, row=r_idx, column=6, value=alt.get('category', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_alerts, row=r_idx, column=7, value=alt.get('description', ''), alignment=Alignment(horizontal='left'))

    # Sheet 5: Devices List
    ws_devs = wb.create_sheet(title='Список ПК')
    ws_devs.views.sheetView[0].showGridLines = True
    dev_headers = ['ID', 'Имя станции', 'Инвентарный №', 'IP-адрес', 'Корпус', 'Этаж', 'Кабинет', 'Группа', 'Статус питания', 'Агент', 'ЦП %', 'ОЗУ %', 'Диск %']
    for idx, h in enumerate(dev_headers, start=1):
        _apply_header_style(ws_devs.cell(row=1, column=idx), h, bg_color='065F46')

    for r_idx, d in enumerate(devices, start=2):
        inv_no = d.get('assetTag', d.get('asset_tag', '')) or '—'
        _set_cell(ws_devs, row=r_idx, column=1, value=d.get('id', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=2, value=d.get('name', ''), alignment=Alignment(horizontal='left'))
        _set_cell(ws_devs, row=r_idx, column=3, value=inv_no, alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=4, value=d.get('ip_address', d.get('ip', '')), alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=5, value=d.get('building', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=6, value=d.get('floor', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=7, value=d.get('room', ''), alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=8, value=d.get('group_name', d.get('group', '')), alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=9, value=str(d.get('power_status', d.get('powerStatus', ''))), alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=10, value=str(d.get('agent_status', d.get('agentStatus', ''))), alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=11, value=d.get('cpu', 0), alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=12, value=d.get('ram', 0), alignment=Alignment(horizontal='center'))
        _set_cell(ws_devs, row=r_idx, column=13, value=d.get('disk', 0), alignment=Alignment(horizontal='center'))

    for ws in [ws_summary, ws_telem, ws_power, ws_alerts, ws_devs]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # On summary sheet, ignore banner and KPI rows to prevent distorting column widths
                if ws == ws_summary and cell.row < chart_table_start_row:
                    continue
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def generate_itilium_excel_report(
    devices: List[Dict[str, Any]],
    hardware_specs: Dict[str, Any],
    scope_title: str = "Весь парк ПК (Fleet)"
) -> bytes:
    """
    Generate an ITIL/CMDB hardware inventory workbook for 1C:Itilium / Service Desk.
    Features dedicated columns for PC name, inventory number, and each physical hardware unit.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Оборудование (Итилиум)"
    ws.views.sheetView[0].showGridLines = True

    # Helper formatters for hardware components
    def _fmt_mb(spec: Dict[str, Any]) -> str:
        mb = spec.get("motherboard") or {}
        if not mb or not isinstance(mb, dict):
            return "—"
        manuf = str(mb.get("manufacturer") or "").strip()
        model = str(mb.get("model") or "").strip()
        serial = str(mb.get("serialNumber") or mb.get("serial") or "").strip()
        ver = str(mb.get("version") or "").strip()
        parts = []
        if manuf and manuf.lower() not in ["oem", "unknown", "to be filled by o.e.m.", "system manufacturer"]:
            parts.append(manuf)
        if model and model.lower() not in ["motherboard", "unknown", "to be filled by o.e.m.", "base board", "system product name"]:
            parts.append(model)
        if not parts:
            parts.append(model or manuf or "—")
        res = " ".join(parts)
        extras = []
        if serial and serial.lower() not in ["unknown", "default string", "none", "to be filled by o.e.m.", "0"]:
            extras.append(f"S/N: {serial}")
        if ver and ver.lower() not in ["unknown", "rev 1.0", "none"]:
            extras.append(f"Rev: {ver}")
        if extras:
            res += f" ({', '.join(extras)})"
        return res if res != "—" else "—"

    def _fmt_cpu(spec: Dict[str, Any], dev: Dict[str, Any]) -> str:
        cpu = spec.get("cpu") or {}
        if not cpu or not isinstance(cpu, dict):
            return "—"
        model = str(cpu.get("model") or "").strip()
        cores = cpu.get("cores")
        threads = cpu.get("threads")
        freq = cpu.get("baseFrequencyGhz") or cpu.get("clockSpeedMhz")
        socket = str(cpu.get("socket") or "").strip()
        if not model:
            return "—"
        details = []
        if cores:
            details.append(f"{cores} ядер")
        if threads and threads != cores:
            details.append(f"{threads} потоков")
        if freq:
            if isinstance(freq, (int, float)) and freq > 100:  # MHz
                details.append(f"{round(freq / 1000, 2)} ГГц")
            elif isinstance(freq, (int, float)):  # GHz
                details.append(f"{freq} ГГц")
            elif str(freq).strip():
                details.append(str(freq).strip())
        if socket and socket.lower() not in ["unknown", "none"]:
            details.append(f"Socket {socket}")
        if details:
            return f"{model} ({', '.join(details)})"
        return model

    def _fmt_ram(spec: Dict[str, Any], dev: Dict[str, Any]) -> str:
        ram = spec.get("ram") or {}
        total_gb = ram.get("totalGb") if isinstance(ram, dict) else None
        if not total_gb or total_gb == 0:
            total_gb = dev.get("ram_total") or dev.get("hardware", {}).get("ram", {}).get("totalGb")
        slots = ram.get("slots") if isinstance(ram, dict) else []
        if not slots and isinstance(dev.get("hardware", {}).get("ram", {}).get("slots"), list):
            slots = dev.get("hardware", {}).get("ram", {}).get("slots", [])
        
        if not total_gb and not slots:
            return "—"
        
        res = f"{total_gb} GB" if total_gb else ""
        if slots and isinstance(slots, list):
            slot_details = []
            for s in slots:
                if not isinstance(s, dict):
                    continue
                cap = s.get("capacityGb") or s.get("sizeGb")
                mfg = str(s.get("manufacturer") or "").strip()
                part = str(s.get("partNumber") or "").strip()
                spd = s.get("speed")
                stype = str(s.get("type") or "").strip()
                sn = str(s.get("serialNumber") or s.get("serial") or "").strip()
                
                mod_str = f"{cap} GB" if cap else ""
                if stype and stype.lower() != "unknown":
                    mod_str += f" {stype}"
                if spd:
                    mod_str += f"-{spd}"
                if mfg and mfg.lower() not in ["unknown", "none"]:
                    mod_str += f" {mfg}"
                if part and part.lower() not in ["unknown", "none"]:
                    mod_str += f" [{part}]"
                if sn and sn.lower() not in ["unknown", "none"]:
                    mod_str += f" (S/N: {sn})"
                if mod_str:
                    slot_details.append(mod_str.strip())
            if slot_details:
                if res:
                    res += f" ({len(slots)} мод: {'; '.join(slot_details)})"
                else:
                    res = f"{'; '.join(slot_details)}"
        return res or "—"

    def _fmt_storage(spec: Dict[str, Any], dev: Dict[str, Any]) -> str:
        disks = spec.get("storage") or spec.get("disks") or []
        if not disks and isinstance(dev.get("drives"), list):
            drv_strs = []
            for drv in dev["drives"]:
                drv_strs.append(f"{drv.get('device', 'Диск')} {drv.get('volumeName', '')} ({drv.get('sizeGb', '')} GB)")
            if drv_strs:
                return "; ".join(drv_strs)
        if not disks or not isinstance(disks, list):
            return "—"
        disk_items = []
        for d in disks:
            if not isinstance(d, dict):
                continue
            model = str(d.get("model") or d.get("name") or "").strip()
            cap = d.get("capacityGb") or d.get("sizeGb")
            mtype = str(d.get("mediaType") or d.get("type") or "").strip()
            bus = str(d.get("busType") or d.get("interfaceType") or "").strip()
            sn = str(d.get("serialNumber") or d.get("serial") or "").strip()
            
            type_tag = mtype.upper() if mtype else ("SSD" if "ssd" in model.lower() else "HDD")
            cap_str = f"{cap} GB" if cap else ""
            header = f"[{type_tag} {cap_str}]".replace("  ", " ").strip()
            
            desc = f"{header} {model}".strip()
            extras = []
            if sn and sn.lower() not in ["unknown", "none", "0"]:
                extras.append(f"S/N: {sn}")
            if bus and bus.lower() not in ["unknown", "none"]:
                extras.append(bus)
            if extras:
                desc += f" ({', '.join(extras)})"
            disk_items.append(desc)
        return ";\n".join(disk_items) if disk_items else "—"

    def _fmt_gpu(spec: Dict[str, Any]) -> str:
        gpus = spec.get("gpus") or spec.get("gpu") or []
        if isinstance(gpus, dict):
            gpus = [gpus]
        if not gpus or not isinstance(gpus, list):
            return "—"
        gpu_items = []
        for g in gpus:
            if not isinstance(g, dict):
                continue
            model = str(g.get("model") or g.get("name") or "").strip()
            vram = g.get("vramGb") or g.get("memoryGb")
            driver = str(g.get("driverVersion") or "").strip()
            if not model or model.lower() in ["unknown", "none"]:
                continue
            item = model
            extras = []
            if vram:
                extras.append(f"{vram} GB VRAM")
            if driver and driver.lower() not in ["unknown", "none"]:
                extras.append(f"Драйвер: {driver}")
            if extras:
                item += f" ({', '.join(extras)})"
            gpu_items.append(item)
        return "; ".join(gpu_items) if gpu_items else "—"

    def _fmt_net(spec: Dict[str, Any], dev: Dict[str, Any]) -> str:
        nets = spec.get("network") or spec.get("netAdapters") or []
        if not nets or not isinstance(nets, list):
            ip = dev.get("ip_address") or dev.get("ip") or "—"
            mac = dev.get("mac_address") or dev.get("mac") or "—"
            return f"MAC: {mac}, IP: {ip}"
        net_items = []
        for n in nets:
            if not isinstance(n, dict):
                continue
            name = str(n.get("name") or n.get("description") or "").strip()
            mac = str(n.get("mac") or n.get("macAddress") or "").strip()
            ip = str(n.get("ip") or n.get("ipAddress") or "").strip()
            spd = n.get("speedMbps")
            if not mac and not ip:
                continue
            item = name if name else "Сетевой адаптер"
            details = []
            if mac:
                details.append(f"MAC: {mac}")
            if ip:
                details.append(f"IP: {ip}")
            if spd:
                details.append(f"{spd} Мбит/с")
            if details:
                item += f" ({', '.join(details)})"
            net_items.append(item)
        return ";\n".join(net_items) if net_items else f"MAC: {dev.get('mac_address', '—')}, IP: {dev.get('ip_address', '—')}"

    def _fmt_os(spec: Dict[str, Any], dev: Dict[str, Any]) -> str:
        os_info = spec.get("os") or {}
        caption = os_info.get("caption") or dev.get("osVersion") or dev.get("os_version") or dev.get("osType") or "Windows"
        arch = os_info.get("architecture") or ("64-bit" if "64" in str(caption) else "")
        build = os_info.get("buildNumber") or os_info.get("build")
        res = str(caption).strip()
        extras = []
        if arch and arch.lower() not in res.lower():
            extras.append(arch)
        if build:
            extras.append(f"Сборка {build}")
        if extras:
            res += f" ({', '.join(extras)})"
        return res

    def _fmt_bios(spec: Dict[str, Any]) -> str:
        bios = spec.get("bios") or {}
        if not bios or not isinstance(bios, dict):
            return "—"
        vendor = str(bios.get("vendor") or "").strip()
        ver = str(bios.get("version") or "").strip()
        date = str(bios.get("releaseDate") or "").strip()
        parts = [p for p in [vendor, ver, date] if p and p.lower() not in ["unknown", "none"]]
        return " ".join(parts) if parts else "—"

    # Row 1: Header Banner
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 10
    total_cols = 16
    end_col_letter = get_column_letter(total_cols)
    
    ws.merge_cells(f"A1:{end_col_letter}1")
    c1 = ws['A1']
    c1.value = _clean_val('ВЫГРУЗКА СПЕЦИФИКАЦИИ ОБОРУДОВАНИЯ ДЛЯ 1С:ИТИЛИУМ / SERVICE DESK')
    c1.font = Font(name='Segoe UI', size=13, bold=True, color='FFFFFF')
    c1.fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
    c1.alignment = Alignment(horizontal='center', vertical='center')

    now_local = datetime.now().astimezone().strftime('%d.%m.%Y %H:%M')
    ws.merge_cells(f"A2:{end_col_letter}2")
    c2 = ws['A2']
    c2.value = _clean_val(f'Зона охвата: {scope_title}   |   Всего станций: {len(devices)}   |   Сформирован: {now_local}')
    c2.font = Font(name='Segoe UI', size=10, italic=False, color='334155')
    c2.fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    c2.alignment = Alignment(horizontal='center', vertical='center')

    # Row 4: Table Headers
    ws.row_dimensions[4].height = 28
    headers = [
        '№',
        'Имя ПК',
        'Инвентарный номер',
        'ID станции',
        'Расположение',
        'Группа / Отдел',
        'Статус',
        'Материнская плата',
        'Процессор (CPU)',
        'Оперативная память (RAM)',
        'Накопители (HDD/SSD)',
        'Видеокарта (GPU)',
        'Сетевой адаптер (MAC / IP)',
        'Операционная система',
        'BIOS',
        'Текущий пользователь'
    ]

    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx)
        _apply_header_style(cell, h_text, bg_color='1E293B', font_color='FFFFFF')

    # Rows 5+: Device Rows
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    fill_alt = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    current_row = 5
    for idx, d in enumerate(devices, start=1):
        dev_id = str(d.get('id', '')).strip()
        spec = hardware_specs.get(dev_id) or hardware_specs.get(dev_id.upper()) or hardware_specs.get(dev_id.lower()) or {}
        row_fill = fill_white if idx % 2 == 1 else fill_alt

        inv_num = d.get('assetTag') or d.get('asset_tag') or '—'
        name_val = d.get('name') or dev_id
        
        b = d.get('building', '') or ''
        f = d.get('floor', '') or ''
        r = d.get('room', '') or ''
        loc_parts = [p for p in [b, f, r] if p]
        location_str = " / ".join(loc_parts) if loc_parts else "—"

        grp = d.get('group_name') or d.get('group') or '—'
        p_stat = str(d.get('power_status', d.get('powerStatus', ''))).lower()
        stat_str = 'В сети' if p_stat in ['on', 'booting'] else 'Выключен'
        user_str = d.get('currentUser') or d.get('current_user') or '—'

        mb_str = _fmt_mb(spec)
        cpu_str = _fmt_cpu(spec, d)
        ram_str = _fmt_ram(spec, d)
        disk_str = _fmt_storage(spec, d)
        gpu_str = _fmt_gpu(spec)
        net_str = _fmt_net(spec, d)
        os_str = _fmt_os(spec, d)
        bios_str = _fmt_bios(spec)

        row_vals = [
            (idx, align_center),
            (name_val, align_left),
            (inv_num, align_center),
            (dev_id, align_center),
            (location_str, align_left),
            (grp, align_left),
            (stat_str, align_center),
            (mb_str, align_left),
            (cpu_str, align_left),
            (ram_str, align_left),
            (disk_str, align_left),
            (gpu_str, align_left),
            (net_str, align_left),
            (os_str, align_left),
            (bios_str, align_left),
            (user_str, align_left)
        ]

        ws.row_dimensions[current_row].height = 24 if ('\n' not in disk_str and '\n' not in net_str) else 38

        for c_idx, (val, alignment) in enumerate(row_vals, start=1):
            c = _set_cell(ws, row=current_row, column=c_idx, value=val, alignment=alignment)
            c.border = thin_border
            c.fill = row_fill
            c.font = Font(name='Segoe UI', size=10)

        current_row += 1

    ws.freeze_panes = 'A5'

    # Auto column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:
                continue
            val_str = str(cell.value or '')
            for line in val_str.split('\n'):
                if len(line) > max_len:
                    max_len = len(line)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 13), 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


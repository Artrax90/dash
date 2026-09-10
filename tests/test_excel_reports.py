import io
import pytest
import openpyxl
from datetime import datetime, timezone, timedelta
from backend.app.services.excel_report_service import generate_monitoring_excel_report

def test_generate_monitoring_excel_report_structure():
    mock_devices = [
        {
            'id': 'PC-01',
            'name': 'ARM-01',
            'group_name': 'Accounting',
            'building': 'Main',
            'floor': 'Floor 2',
            'room': '204',
            'power_status': 'On',
            'agent_status': 'Connected',
            'ip_address': '192.168.1.50',
            'cpu': 25,
            'ram': 45,
            'disk': 60,
            'assetTag': 'INV-2026-001'
        },
        {
            'id': 'PC-02',
            'name': 'ARM-02',
            'group_name': 'Accounting',
            'building': 'Main',
            'floor': 'Floor 2',
            'room': '204',
            'power_status': 'Off',
            'agent_status': 'Disconnected',
            'ip_address': '192.168.1.51',
            'cpu': 0,
            'ram': 0,
            'disk': 70,
            'assetTag': 'INV-2026-002'
        }
    ]
    
    mock_telemetry_points = [
        {
            'time': (datetime.now(timezone.utc) - timedelta(hours=i)).timestamp(),
            'timestamp': (datetime.now(timezone.utc) - timedelta(hours=i)).isoformat(),
            'deviceId': 'PC-01',
            'deviceName': 'ARM-01',
            'cpu': 20 + i * 2,
            'ram': 40 + i,
            'disk': 60,
            'isOnline': True
        }
        for i in range(12)
    ]
    
    mock_power_events = [
        {
            'id': 'EVT-1',
            'timestamp': datetime.now(timezone.utc).isoformat(),
            'action': 'WAKE',
            'title': 'Wake by Schedule (WoL)',
            'deviceId': 'PC-01',
            'deviceName': 'ARM-01',
            'status': 'Success',
            'initiator': 'Schedule',
            'source': 'SCHEDULE',
            'details': 'Morning Wake'
        }
    ]
    
    mock_alerts = [
        {
            'id': 'ALT-1',
            'timestamp': datetime.now(timezone.utc).isoformat(),
            'deviceId': 'PC-01',
            'severity': 'Warning',
            'description': 'High CPU utilization (95 pct)',
            'category': 'Resource'
        }
    ]

    report_bytes = generate_monitoring_excel_report(
        period_type='hourly_24h',
        scope_title='Accounting (Main / Floor 2 / 204)',
        devices=mock_devices,
        telemetry_points=mock_telemetry_points,
        power_events=mock_power_events,
        alerts=mock_alerts
    )

    assert isinstance(report_bytes, bytes)
    assert len(report_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(report_bytes))
    sheet_names = wb.sheetnames
    
    assert 'Сводка и Графики' in sheet_names
    assert 'Телеметрия' in sheet_names
    assert 'События питания' in sheet_names
    assert 'Алерты и Инциденты' in sheet_names
    assert 'Список ПК' in sheet_names

    summary_sheet = wb['Сводка и Графики']
    assert len(summary_sheet._charts) >= 1
    # Check that A2 does not have "UTC" and is formatted with local time
    assert "UTC" not in summary_sheet['A2'].value
    # Check that the 12 rows have distinct time intervals
    chart_row_times = [summary_sheet.cell(r, 1).value for r in range(10, 22)]
    assert len(set(chart_row_times)) == 12

    # Check Sheet 2 (Телеметрия): Col 2 is "ID ПК", Col 3 is "Имя ПК"
    telem_sheet = wb['Телеметрия']
    assert telem_sheet.cell(1, 2).value == 'ID ПК'
    assert telem_sheet.cell(1, 3).value == 'Имя ПК'
    assert telem_sheet.cell(2, 2).value == 'PC-01'
    assert telem_sheet.cell(2, 3).value == 'ARM-01'

    # Check Sheet 3 (События питания): Col 5 is "ID ПК", Col 6 is "Имя ПК"
    power_sheet = wb['События питания']
    assert power_sheet.cell(1, 5).value == 'ID ПК'
    assert power_sheet.cell(1, 6).value == 'Имя ПК'
    assert power_sheet.cell(2, 5).value == 'PC-01'
    assert power_sheet.cell(2, 6).value == 'ARM-01'

    # Check Sheet 4 (Алерты и Инциденты): Col 3 is "ID ПК", Col 4 is "Имя ПК"
    alert_sheet = wb['Алерты и Инциденты']
    assert alert_sheet.cell(1, 3).value == 'ID ПК'
    assert alert_sheet.cell(1, 4).value == 'Имя ПК'
    assert alert_sheet.cell(2, 3).value == 'PC-01'
    assert alert_sheet.cell(2, 4).value == 'ARM-01'

    # Check Sheet 5 (Список ПК): Col 3 is "Инвентарный №"
    dev_sheet = wb['Список ПК']
    assert dev_sheet.cell(1, 3).value == 'Инвентарный №'
    assert dev_sheet.cell(2, 3).value == 'INV-2026-001'

@pytest.mark.anyio
async def test_excel_report_endpoint():
    import httpx
    from backend.app.main import app
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        resp = await client.get("/api/v1/devices/reports/excel?time_range=24h&group=ALL")
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert len(resp.content) > 0
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert "Сводка и Графики" in wb.sheetnames

def test_generate_monitoring_excel_report_with_illegal_xml_control_characters():
    """TDD: Ensure excel report generation does NOT crash on control characters like \x00, \x05, etc."""
    mock_devices = [
        {
            'id': 'PC-9AEB',
            'name': 'ARM-09\x00\x01',
            'group_name': 'Auditorium\x07',
            'building': 'Main\x08',
            'floor': 'Floor 1',
            'room': '101\x1f',
            'power_status': 'On',
            'agent_status': 'Connected',
            'ip_address': '172.16.42.0',
            'cpu': 10,
            'ram': 20,
            'disk': 30
        }
    ]
    mock_power_events = [
        {
            'id': 'EVT-USB-1',
            'timestamp': datetime.now(timezone.utc).isoformat(),
            'action': 'USB_PLUG',
            'title': 'USB Device Attached\x00',
            'deviceId': 'PC-9AEB',
            'deviceName': 'ARM-09\x05',
            'status': 'Success',
            'initiator': 'Agent\x00',
            'source': 'AGENT',
            'details': 'Flash drive inserted\x02'
        }
    ]
    mock_alerts = [
        {
            'id': 'ALT-USB-99',
            'timestamp': datetime.now(timezone.utc).isoformat(),
            'deviceId': 'PC-9AEB',
            'severity': 'Info',
            'description': 'Извлечен съемный USB-накопитель: Netac OnlyDisk USB Device (S/N: 0000000005\x00)',
            'category': 'Hardware\x00'
        }
    ]
    mock_telemetry_points = [
        {
            'time': datetime.now(timezone.utc).timestamp(),
            'timestamp': datetime.now(timezone.utc).isoformat(),
            'deviceId': 'PC-9AEB',
            'cpu': 15,
            'ram': 35,
            'disk': 45,
            'isOnline': True
        }
    ]

    report_bytes = generate_monitoring_excel_report(
        period_type='hourly_24h',
        scope_title='Test Scope with \x00 null byte',
        devices=mock_devices,
        telemetry_points=mock_telemetry_points,
        power_events=mock_power_events,
        alerts=mock_alerts
    )

    assert isinstance(report_bytes, bytes)
    assert len(report_bytes) > 0
    wb = openpyxl.load_workbook(io.BytesIO(report_bytes))
    ws_alerts = wb['Алерты и Инциденты']
    # Verify description is saved and illegal char was stripped (Column 7 is description now that Device Name is Column 4)
    assert 'Netac OnlyDisk USB Device' in ws_alerts.cell(row=2, column=7).value
    assert '\x00' not in ws_alerts.cell(row=2, column=7).value

def test_generate_itilium_excel_report():
    from backend.app.services.excel_report_service import generate_itilium_excel_report

    mock_devices = [
        {
            'id': 'PC-01',
            'name': 'ARM-01',
            'group_name': 'Бухгалтерия',
            'building': 'Главный корпус',
            'floor': '2 этаж',
            'room': 'Каб. 204',
            'power_status': 'On',
            'agent_status': 'Connected',
            'ip_address': '192.168.1.50',
            'mac_address': '70:85:C2:54:1F:7D',
            'assetTag': 'INV-2026-ITIL-001',
            'currentUser': 'Иванов И.И.'
        }
    ]

    mock_specs = {
        'PC-01': {
            'motherboard': {
                'manufacturer': 'ASUSTeK COMPUTER INC.',
                'model': 'PRIME B450M-K',
                'serialNumber': '21038472910'
            },
            'bios': {
                'vendor': 'American Megatrends Inc.',
                'version': '2.40',
                'releaseDate': '12.05.2023'
            },
            'cpu': {
                'model': 'AMD Ryzen 5 3600',
                'cores': 6,
                'threads': 12,
                'baseFrequencyGhz': 3.6,
                'socket': 'AM4'
            },
            'ram': {
                'totalGb': 16,
                'slots': [
                    {'slot': 'DIMM 0', 'capacityGb': 8, 'speed': 3200, 'manufacturer': 'Kingston', 'partNumber': 'KHX3200C16D4/8GX', 'serialNumber': '98A72B10'},
                    {'slot': 'DIMM 1', 'capacityGb': 8, 'speed': 3200, 'manufacturer': 'Kingston', 'partNumber': 'KHX3200C16D4/8GX', 'serialNumber': '98A72B11'}
                ]
            },
            'storage': [
                {'model': 'Kingston SA400S37480G', 'capacityGb': 480, 'serialNumber': '50026B7683C01234', 'mediaType': 'SSD', 'busType': 'SATA'},
                {'model': 'WDC WD10EZEX-08WN4A0', 'capacityGb': 1000, 'serialNumber': 'WD-WCC6Y7890123', 'mediaType': 'HDD', 'busType': 'SATA'}
            ],
            'gpus': [
                {'model': 'NVIDIA GeForce GTX 1660 SUPER', 'vramGb': 6, 'driverVersion': '536.23'}
            ],
            'network': [
                {'name': 'Realtek PCIe GbE Family Controller', 'mac': '70:85:C2:54:1F:7D', 'ip': '192.168.1.50', 'speedMbps': 1000}
            ]
        }
    }

    report_bytes = generate_itilium_excel_report(
        devices=mock_devices,
        hardware_specs=mock_specs,
        scope_title='Бухгалтерия (Главный корпус / 2 этаж / Каб. 204)'
    )

    assert isinstance(report_bytes, bytes)
    assert len(report_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(report_bytes))
    sheet = wb.active
    assert 'Итилиум' in sheet.title or 'Оборудование' in sheet.title

    # Header check in row 1 (no banner, direct 1C/Itilium import format)
    headers = [sheet.cell(1, c).value for c in range(1, 20) if sheet.cell(1, c).value is not None]
    assert 'Имя ПК' in headers
    assert 'Штрих-код КЕ' in headers
    assert 'Инвентарный номер' not in headers
    assert 'Материнская плата' in headers
    assert 'Процессор (CPU)' in headers
    assert 'Оперативная память (RAM)' in headers
    assert 'Накопители (HDD/SSD)' in headers
    assert 'Видеокарта (GPU)' in headers

    # Ensure unwanted columns are removed
    assert len(headers) == 13
    assert 'ID станции' not in headers
    assert 'Статус' not in headers
    assert 'Статус питания' not in headers
    assert 'Текущий пользователь' not in headers

    # Value checks in row 2 (data starts immediately after row 1 headers)
    name_col = headers.index('Имя ПК') + 1
    inv_col = headers.index('Штрих-код КЕ') + 1
    mb_col = headers.index('Материнская плата') + 1
    cpu_col = headers.index('Процессор (CPU)') + 1
    ram_col = headers.index('Оперативная память (RAM)') + 1
    disk_col = headers.index('Накопители (HDD/SSD)') + 1
    gpu_col = headers.index('Видеокарта (GPU)') + 1

    assert sheet.cell(2, name_col).value == 'ARM-01'
    assert sheet.cell(2, inv_col).value == 'INV-2026-ITIL-001'
    assert 'PRIME B450M-K' in str(sheet.cell(2, mb_col).value)
    assert 'Ryzen 5 3600' in str(sheet.cell(2, cpu_col).value)
    assert '16 GB' in str(sheet.cell(2, ram_col).value)
    assert 'Kingston SA400' in str(sheet.cell(2, disk_col).value)
    assert 'GTX 1660' in str(sheet.cell(2, gpu_col).value)
    assert sheet.freeze_panes == 'A2'

@pytest.mark.anyio
async def test_itilium_report_endpoint():
    import httpx
    from backend.app.main import app
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        resp = await client.get("/api/v1/devices/reports/itilium")
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert len(resp.content) > 0
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        sheet = wb.active
        assert any('Итилиум' in s or 'Оборудование' in s for s in wb.sheetnames)
        assert sheet.cell(1, 1).value == '№'
        headers = [sheet.cell(1, c).value for c in range(1, 20) if sheet.cell(1, c).value is not None]
        assert 'Штрих-код КЕ' in headers


